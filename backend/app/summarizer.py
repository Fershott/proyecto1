import io
import re
from collections import Counter
"""Motor de resúmenes y extracción de texto utilizado por CogniCore."""

from typing import Any, List, TYPE_CHECKING

try:  # pragma: no cover - optional dependency for API runtime
    from fastapi import HTTPException  # type: ignore
except ModuleNotFoundError:  # pragma: no cover - offline execution fallback
    class HTTPException(Exception):
        """Ligera implementación para uso fuera de FastAPI."""

        def __init__(self, status_code: int, detail: str):
            self.status_code = status_code
            self.detail = detail
            super().__init__(detail)


if TYPE_CHECKING:  # pragma: no cover - solo para tips de tipos
    from fastapi import UploadFile  # type: ignore
else:
    UploadFile = Any


def _extract_text_from_pdf(file: "UploadFile") -> str:
    """Extrae texto desde archivos PDF página por página."""
    try:
        from PyPDF2 import PdfReader  # type: ignore
    except ModuleNotFoundError as exc:  # pragma: no cover - dependencia opcional
        raise HTTPException(status_code=500, detail="PyPDF2 no está disponible") from exc

    buffer = io.BytesIO(file.file.read())
    reader = PdfReader(buffer)
    text_parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(text_parts)


def _extract_text_from_docx(file: "UploadFile") -> str:
    """Lee el contenido de documentos .docx y concatena sus párrafos."""
    try:
        from docx import Document  # type: ignore
    except ModuleNotFoundError as exc:  # pragma: no cover - dependencia opcional
        raise HTTPException(status_code=500, detail="python-docx no está disponible") from exc

    buffer = io.BytesIO(file.file.read())
    document = Document(buffer)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_text_from_pptx(file: "UploadFile") -> str:
    """Recorre las diapositivas de una presentación y reúne el texto visible."""
    try:
        from pptx import Presentation  # type: ignore
    except ModuleNotFoundError as exc:  # pragma: no cover - dependencia opcional
        raise HTTPException(status_code=500, detail="python-pptx no está disponible") from exc

    buffer = io.BytesIO(file.file.read())
    presentation = Presentation(buffer)
    text_runs: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_runs.append(shape.text)
    return "\n".join(text_runs)


def _extract_text_from_plain(file: UploadFile) -> str:
    """Decodifica archivos de texto plano en UTF-8 ignorando errores."""
    return file.file.read().decode("utf-8", errors="ignore")


EXTRACTORS = {
    "application/pdf": _extract_text_from_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_text_from_docx,
    "application/msword": _extract_text_from_docx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _extract_text_from_pptx,
    "text/plain": _extract_text_from_plain,
}


STOPWORDS = {
    "el",
    "la",
    "los",
    "las",
    "de",
    "y",
    "que",
    "en",
    "a",
    "un",
    "una",
    "es",
    "para",
    "con",
    "por",
    "del",
    "se",
    "al",
    "lo",
    "como",
    "más",
    "su",
    "sus",
    "no",
}


def _tokenize_sentences(text: str) -> List[str]:
    """Separa un texto en oraciones usando puntuación como delimitador."""
    cleaned = re.sub(r"\s+", " ", text.strip())
    if not cleaned:
        return []
    sentences = re.split(r"(?<=[.!?])\s+", cleaned)
    return [sentence.strip() for sentence in sentences if sentence.strip()]


def _score_sentence(sentence: str, word_frequencies: Counter) -> float:
    """Calcula la relevancia de una oración según la frecuencia de sus palabras."""
    words = re.findall(r"[\wáéíóúñ]+", sentence.lower())
    score = sum(word_frequencies[word] for word in words if word not in STOPWORDS)
    return score / (len(words) or 1)


def summarize_text(text: str, sentences: int = 5) -> List[str]:
    """Resume el texto seleccionando las oraciones mejor puntuadas."""
    segments = _tokenize_sentences(text)
    if not segments:
        return []

    desired = max(1, min(sentences, len(segments)))
    words = re.findall(r"[\wáéíóúñ]+", text.lower())
    frequencies = Counter(word for word in words if word not in STOPWORDS)
    if not frequencies:
        return segments[:desired]

    scored = [(segment, _score_sentence(segment, frequencies)) for segment in segments]
    top_segments = sorted(scored, key=lambda item: item[1], reverse=True)[:desired]
    ordered = sorted(top_segments, key=lambda item: segments.index(item[0]))
    return [segment for segment, _ in ordered]


def extract_keywords(text: str, limit: int = 10) -> List[str]:
    """Obtiene las palabras clave más frecuentes excluyendo stopwords."""
    words = re.findall(r"[\wáéíóúñ]{4,}", text.lower())
    frequencies = Counter(word for word in words if word not in STOPWORDS)
    most_common = frequencies.most_common(limit)
    return [word for word, _ in most_common]


async def generate_summary_from_file(file: UploadFile, sentences: int = 5) -> tuple[str, str]:
    """Genera un resumen desde un archivo soportado y devuelve también el texto original."""
    extractor = EXTRACTORS.get(file.content_type)
    if not extractor:
        raise HTTPException(status_code=400, detail="Tipo de archivo no soportado")

    try:
        text = extractor(file)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"No se pudo procesar el archivo: {exc}") from exc
    finally:
        file.file.close()

    summary_segments = summarize_text(text, sentences)
    if not summary_segments:
        fallback_segments = _tokenize_sentences(text)[:max(1, sentences)]
        summary_segments = fallback_segments if fallback_segments else [text.strip()]
    return " ".join(segment for segment in summary_segments if segment), text


async def summarize_from_text(text: str, sentences: int = 5) -> tuple[str, str]:
    """Crea un resumen a partir de texto crudo enviado por el cliente."""
    summary_segments = summarize_text(text, sentences)
    if not summary_segments:
        fallback_segments = _tokenize_sentences(text)[:max(1, sentences)]
        summary_segments = fallback_segments if fallback_segments else [text.strip()]
    return " ".join(segment for segment in summary_segments if segment), text


async def generate_summary(file: UploadFile | None, text: str | None, sentences: int) -> tuple[str, str]:
    """Determina si se debe resumir un archivo o texto directo."""
    if file is None and not text:
        raise HTTPException(status_code=400, detail="Debe proporcionar un archivo o texto para resumir")

    if file is not None:
        return await generate_summary_from_file(file, sentences)

    return await summarize_from_text(text or "", sentences)


def top_keywords(text: str) -> List[str]:
    """Expone un helper para recuperar palabras clave relevantes."""
    return extract_keywords(text)
